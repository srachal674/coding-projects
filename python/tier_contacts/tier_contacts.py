import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import openpyxl
import sqlite3
from pathlib import Path

# DATABASE SETUP
# Creates the SQLite tables used to store students and their course enrollments.
def setup_database():
    connection = sqlite3.connect("tier_contacts.db")
    cursor = connection.cursor()

    # Store one record per student. Student ID is the unique key.
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS students (
            student_id TEXT PRIMARY KEY, 
            name TEXT NOT NULL,
            grade TEXT,
            phone TEXT
        )
    """)

    # Store each student's course enrollments separately so one student can have multiple courses.
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS enrollments (
            student_id TEXT NOT NULL,
            course TEXT NOT NULL,
            PRIMARY KEY (student_id, course),
            FOREIGN KEY(student_id) REFERENCES students(student_id)
        )
    """)

    connection.commit()
    connection.close()

# ROSTER IMPORT
# Prompts the user to select a PowerSchool roster spreadsheet and imports its data.
def import_roster(roster_table):
    file_path = filedialog.askopenfilename(
        title="Select PowerSchool Roster",
        filetypes=[
            ("Excel files", "*.xlsx"),
            ("All files", "*.*")
        ]
    )

    # Only continue if the user selected a file.
    if file_path:
        # Open the selected Excel workbook and use its active worksheet.
        workbook = openpyxl.load_workbook(file_path)
        print("Workbook opened successfully!")
        print("Worksheets:", workbook.sheetnames)
        worksheet = workbook.active

        # Read the spreadsheet headers so columns can be found by name instead of fixed position.
        headers = [cell.value for cell in worksheet[1]]
        required_headers = ['Id', 'Name', 'Grade', 'Phone', 'Course']
        missing_headers = []

        for header in required_headers:
            if header not in headers:
                missing_headers.append(header)

        # Stop the import if any required PowerSchool columns are missing.
        if missing_headers:
            messagebox.showerror(
                "Invalid Roster",
                "Missing required columns: " + ", ".join(missing_headers)
            )
        else:
            messagebox.showinfo(
                "Roster Valid",
                "The roster contains all required columns."
            )

            # Map each required header to its actual column position in the spreadsheet.
            column_map = {
                "Id": headers.index("Id"),
                "Name": headers.index("Name"),
                "Grade": headers.index("Grade"),
                "Phone": headers.index("Phone"),
                "Course": headers.index("Course")
            }

            # Build a dictionary with one record per Student ID.
            # A student can appear on multiple spreadsheet rows because each course is a separate enrollment.
            students = {}

            for row in worksheet.iter_rows(min_row=2, values_only=True):
                student_id = row[column_map["Id"]]
                name = row[column_map["Name"]]
                grade = row[column_map["Grade"]]
                phone = row[column_map["Phone"]]
                course = row[column_map["Course"]]

                # Create the student only the first time the Student ID appears.
                if student_id not in students:
                    students[student_id] = {
                        "id": student_id,
                        "name": name,
                        "grade": grade,
                        "phone": phone,
                        "courses": []
                    }

                # Add this row's course to the student's course list.
                students[student_id]["courses"].append(course)

            # Save the imported roster to the SQLite database.
            connection = sqlite3.connect("tier_contacts.db")
            cursor = connection.cursor()

            # Save or update each student's basic information.
            for student in students.values():
                cursor.execute("""
                    INSERT OR REPLACE INTO students (
                        student_id,
                        name,
                        grade,
                        phone
                    )
                    VALUES (?, ?, ?, ?)
                """, (
                    student["id"],
                    student["name"],
                    student["grade"],
                    student["phone"]
                ))

                # Save each course enrollment without creating duplicate student/course pairs.
                for course in student["courses"]:
                    cursor.execute("""
                        INSERT OR IGNORE INTO enrollments (
                            student_id,
                            course
                        )
                        VALUES (?, ?)
                    """, (
                        student["id"],
                        course
                    ))

            # Make the database changes permanent, then close the connection.
            connection.commit()
            connection.close()

            # Diagnostic count used to confirm that multi-course students were grouped correctly.
            multiple_course_students = 0

            for student in students.values():
                if len(student["courses"]) > 1:
                    multiple_course_students += 1

            print("Roster saved to database.")
            print("Roster rows:", worksheet.max_row - 1)
            print("Unique students:", len(students))
            print("Students with multiple courses:", multiple_course_students)
            load_roster_table(roster_table)

        print("Headers found:", headers)
        print("Missing headers:", missing_headers)

    else:
        print("No file selected.")

#ROSTER DISPLAY
# Loads the saved student roster from the database into the roster table.
def load_roster_table(roster_table):
    for item in roster_table.get_children():
        roster_table.delete(item)

    connection = sqlite3.connect("tier_contacts.db")
    cursor = connection.cursor()
    cursor.execute("""
        SELECT student_id, name, grade, phone
        FROM students
        ORDER BY name
    """)

    students = cursor.fetchall()

    # Track the longest course list so the Course(s) column can be sized automatically.
    longest_courses = 0

    for student in students:
        student_id = student[0]
        name = student[1]
        grade = student[2]
        phone= student[3]

        cursor.execute("""
            SELECT course
            FROM enrollments
            WHERE student_id = ?
        """, (student_id,))

        course_rows = cursor.fetchall()
        courses = ", ".join(course[0] for course in course_rows)

        # Keep the largest character count found while working through the roster.
        if len(courses) > longest_courses:
            longest_courses = len(courses)

        roster_table.insert(
            "",
            "end",
            values=(student_id, name, grade, phone, courses)
        )

    # Convert the longest character count to an approximate pixel width for the Treeview column.
    roster_table.column("Courses", width=longest_courses * 7, stretch=False)
    connection.close()

# STUDENT ROSTER WINDOW
# Opens the roster screen and builds its controls and table.
def open_roster():
    roster_window = tk.Toplevel(root)

    roster_window.title("Student Roster")
    roster_window.geometry("800x500")
    roster_title = ttk.Label(
        roster_window,
        text="Student Roster",
        font=("Arial", 20, "bold")
    )

    roster_title.pack(pady=20)
    roster_frame = ttk.Frame(roster_window)
    roster_frame.pack(padx=20, pady=10, fill="both", expand=True)

    # Button used to start a new PowerSchool roster import.
    import_button = ttk.Button(
        roster_frame,
        text="Import Roster",
        command=lambda: import_roster(roster_table)
    )
    import_button.pack(pady=10)

    # Table that will display the stored student roster.
    roster_table = ttk.Treeview(
        roster_frame,
        columns=("Id", "Name", "Grade", "Phone", "Courses"),
        show="headings"
    )
    roster_table.heading("Id", text="Id")
    roster_table.heading("Name", text="Name")
    roster_table.heading("Grade", text="Grade")
    roster_table.heading("Phone", text="Phone")
    roster_table.heading("Courses", text="Course(s)")
    roster_table.column("Courses", width=500, stretch=False)

    # Connect a horizontal scrollbar to the roster table for long course lists.
    horizontal_scrollbar = ttk.Scrollbar(
        roster_frame,
        orient="horizontal",
        command=roster_table.xview
    )

    # Connect a vertical scrollbar to the roster table for long student lists.
    vertical_scrollbar = ttk.Scrollbar(
        roster_frame,
        orient="vertical",
        command=roster_table.yview
    )

    roster_table.configure(
        xscrollcommand=horizontal_scrollbar.set,
        yscrollcommand=vertical_scrollbar.set
    )
    horizontal_scrollbar.pack(fill="x")
    vertical_scrollbar.pack(side="right", fill="y")
    
    roster_table.pack(fill="both", expand=True, pady=10)
    load_roster_table(roster_table)
 

# MAIN APPLICATION
# Make sure the database exists before creating the main application window.
setup_database()

root = tk.Tk()
root.title("Tier Contacts")
root.geometry("900x600")

title_label = ttk.Label(
    root,
    text="Tier Contacts",
    font=("Arial", 24, "bold")
)
title_label.pack(pady=30)

# Frame that holds the main navigation buttons.
menu_frame = ttk.Frame(root)
menu_frame.pack(pady=20)

roster_button = ttk.Button(
    menu_frame,
    text="Student Roster",
    command=open_roster
)
roster_button.pack(pady=10)















root.mainloop()